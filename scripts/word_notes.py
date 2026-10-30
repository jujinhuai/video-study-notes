"""Reusable study-note DOCX authoring; run with bundled Python -B.

Import and block construction perform no file I/O. NotesDocument.save() is the
only writer. Content selection, source verification, and rendering stay with
the caller. See references/word-tool.md for the content-draft contract.
"""

from copy import deepcopy
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


@dataclass(frozen=True)
class Block:
    kind: str
    args: tuple
    options: dict = field(default_factory=dict)


@dataclass(frozen=True)
class Source:
    label: str
    url: str = ""
    detail: str = ""


def paragraph(text, style="Normal", emphasis=False):
    """Optionally color explicit bold spans; content priority is caller-selected."""
    return Block("paragraph", (text,), {"style": style, "emphasis": emphasis})


def heading(text, level=1):
    if level not in (1, 2):
        raise ValueError("Study-note headings support levels 1 and 2")
    return Block("heading", (text,), {"level": level})


def equation(*items):
    return Block("equation", items)


def table(headers, rows, widths_cm, *, allow_row_split=False):
    return Block("table", (headers, rows, widths_cm), {"allow_row_split": allow_row_split})


def image(path, caption, alt_text="", width_cm=None):
    return Block("image", (path, caption), {"alt_text": alt_text, "width_cm": width_cm})


def code(text):
    return Block("paragraph", (text,), {"style": "代码"})


def source(text):
    return Block("paragraph", (text,), {"style": "来源"})


def toc():
    """Insert an editable field; check/update its displayed result after rendering."""
    return Block("toc", ())


def _font(target, name, size, bold=False):
    target.font.name = name
    target.font.size = Pt(size)
    target.font.bold = bold
    target.font.color.rgb = RGBColor(0, 0, 0)
    props = target.element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    fonts.attrib.clear()  # Theme fonts otherwise override explicit Chinese fonts.
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), name)


def _indent(style, chars=0):
    props = style.element.get_or_add_pPr()
    indent = props.find(qn("w:ind"))
    if indent is None:
        indent = OxmlElement("w:ind")
        props.append(indent)
    indent.set(qn("w:firstLineChars"), str(chars * 100))
    indent.set(qn("w:firstLine"), "0")


def _math_run(value, plain=False):
    run = OxmlElement("m:r")
    if plain:
        props = OxmlElement("m:rPr")
        style = OxmlElement("m:sty")
        style.set(qn("m:val"), "p")
        props.append(style)
        run.append(props)
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), "Cambria Math")
    props.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    props.append(size)
    run.append(props)
    text = OxmlElement("m:t")
    text.set(qn("xml:space"), "preserve")
    text.text = str(value)
    run.append(text)
    return run


def _math_append(parent, items):
    for item in items if isinstance(items, (tuple, list)) else (items,):
        if isinstance(item, (tuple, list)):
            _math_append(parent, item)
        elif isinstance(item, etree._Element):
            # OOXML append moves nodes; copy to keep batches and total drafts reusable.
            parent.append(deepcopy(item))
        else:
            parent.append(_math_run(item))


def _argument(parent, name, items):
    element = OxmlElement("m:" + name)
    _math_append(element, items)
    parent.append(element)


def _structure(tag, **arguments):
    element = OxmlElement("m:" + tag)
    for name, items in arguments.items():
        _argument(element, name, items)
    return element


def sub(base, value):
    return _structure("sSub", e=base, sub=value)


def sup(base, value):
    return _structure("sSup", e=base, sup=value)


def frac(numerator, denominator):
    return _structure("f", num=numerator, den=denominator)


def bounded(content, opening="(", closing=")"):
    element = OxmlElement("m:d")
    props = OxmlElement("m:dPr")
    for key, value in (("begChr", opening), ("endChr", closing), ("grow", "1")):
        prop = OxmlElement("m:" + key)
        prop.set(qn("m:val"), value)
        props.append(prop)
    element.append(props)
    _argument(element, "e", content)
    return element


def _matrix(rows):
    if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError("Matrix rows must be nonempty and rectangular")
    element = OxmlElement("m:m")
    props = OxmlElement("m:mPr")
    columns = OxmlElement("m:mcs")
    column = OxmlElement("m:mc")
    col_props = OxmlElement("m:mcPr")
    for key, value in (("count", str(len(rows[0]))), ("mcJc", "center")):
        prop = OxmlElement("m:" + key)
        prop.set(qn("m:val"), value)
        col_props.append(prop)
    column.append(col_props)
    columns.append(column)
    props.append(columns)
    element.append(props)
    for values in rows:
        row = OxmlElement("m:mr")
        for value in values:
            _argument(row, "e", value)
        element.append(row)
    return element


def matrix(rows):
    return bounded(_matrix(rows), "[", "]")


def column_vector(values):
    """Each value, including a composite expression, occupies one matrix cell."""
    return matrix([[value] for value in values])


def cases(rows):
    if any(len(row) != 2 for row in rows):
        raise ValueError("Each case needs an expression and condition")
    return bounded(_matrix(rows), "{", "")


def summation(lower, upper, expression):
    element = OxmlElement("m:nary")
    props = OxmlElement("m:naryPr")
    for key, value in (("chr", "∑"), ("limLoc", "undOvr")):
        prop = OxmlElement("m:" + key)
        prop.set(qn("m:val"), value)
        props.append(prop)
    # Absent limits need an explicit hide flag or Word displays an empty box.
    for name, value in (("sub", lower), ("sup", upper)):
        if value is None or (isinstance(value, str) and value == ""):
            hidden = OxmlElement("m:" + name + "Hide")
            hidden.set(qn("m:val"), "1")
            props.append(hidden)
    element.append(props)
    for name, value in (("sub", lower), ("sup", upper), ("e", expression)):
        _argument(element, name, () if name != "e" and value is None else value)
    return element


def function(name, expression):
    element = OxmlElement("m:func")
    label = OxmlElement("m:fName")
    label.append(_math_run(name, plain=True))
    element.append(label)
    _argument(element, "e", bounded(expression))
    return element


class NotesDocument:
    """Build from explicit blocks and sources; never reads earlier deliverables."""

    def __init__(self, title, output_path, *, blocks=(), sources=(), chinese_font="等线"):
        self.title = title
        self.output_path = Path(output_path)
        # Naming defaults belong to the caller; explicit user names may differ.
        if not isinstance(title, str) or not title.strip():
            raise ValueError("Study-note title must be a nonempty string")
        if self.output_path.suffix.lower() != ".docx":
            raise ValueError("Output must be a .docx file")
        if not isinstance(chinese_font, str) or not chinese_font.strip():
            raise ValueError("Chinese font must be a nonempty font name")
        self.chinese_font = chinese_font
        self.doc = Document()
        self._styles()
        self.doc.core_properties.title = title
        self.doc.core_properties.author = ""
        self.doc.add_paragraph(title, "Title")
        self._after_table = False
        for block in blocks:
            self._add(block)
        sources = tuple(sources)
        if sources:
            self._add(heading("来源与范围"))
            for item in sources:
                self._source(item)

    def _styles(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        for edge in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, edge, Cm(2.5))
        left, center, justify = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY
        specs = (
            ("Normal", 11, 1.5, 0, 6, justify, 2, False, self.chinese_font),
            ("Title", 22, 1.2, 0, 18, left, 0, True, self.chinese_font),
            ("Heading 1", 16, 1.2, 18, 9, left, 0, True, self.chinese_font),
            ("Heading 2", 13, 1.2, 12, 6, left, 0, True, self.chinese_font),
            ("来源", 9.5, 1.2, 0, 6, left, 0, False, self.chinese_font),
            ("提示", 11, 1.5, 0, 6, justify, 0, False, self.chinese_font),
            ("公式", 12, 1.0, 8, 8, center, 0, False, "Cambria Math"),
            ("图片", 11, 1.2, 6, 3, center, 0, False, self.chinese_font),
            ("图注", 9.5, 1.2, 3, 12, center, 0, False, self.chinese_font),
            ("表格正文", 10.5, 1.2, 0, 3, left, 0, False, self.chinese_font),
            ("代码", 10.5, 1.15, 0, 6, left, 0, False, "Consolas"),
            ("目录标题", 16, 1.2, 18, 9, left, 0, True, self.chinese_font),
        )
        for name, size, spacing, before, after, align, indent, bold, font in specs:
            style = self.doc.styles[name] if name in self.doc.styles else self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            if name != "Normal":
                style.base_style = self.doc.styles["Normal"]
            _font(style, font, size, bold)
            _indent(style, indent)
            fmt = style.paragraph_format
            fmt.line_spacing = spacing
            fmt.space_before, fmt.space_after = Pt(before), Pt(after)
            fmt.alignment = align
            fmt.keep_with_next = name in ("Title", "Heading 1", "Heading 2", "图片", "目录标题")
            fmt.keep_together = False
            fmt.widow_control = True
            props = style.element.get_or_add_pPr()
            for border in list(props.findall(qn("w:pBdr"))):
                props.remove(border)
            if name in ("Heading 1", "Heading 2"):
                outline = props.find(qn("w:outlineLvl"))
                if outline is None:
                    outline = OxmlElement("w:outlineLvl")
                    props.append(outline)
                outline.set(qn("w:val"), str(int(name[-1]) - 1))
        for level in (1, 2):
            # Word recognizes built-in TOC names in lowercase; custom names collide.
            name = "toc " + str(level)
            style = self.doc.styles[name] if name in self.doc.styles else self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
            style.style_id = "TOC" + str(level)
            style.base_style = self.doc.styles["来源"]
            _font(style, self.chinese_font, 9.5)
            _indent(style)
            fmt = style.paragraph_format
            fmt.line_spacing = 1.1
            fmt.space_before, fmt.space_after = Pt(0), Pt(2)
            fmt.left_indent = Cm(0 if level == 1 else 0.6)
            fmt.keep_with_next = False

    def _paragraph(self, text, style, emphasis=False):
        paragraph = self.doc.add_paragraph(style=style)
        if self._after_table:
            paragraph.paragraph_format.space_before = Pt(8)
        if style == "代码":
            paragraph.add_run(str(text))  # Preserve code whitespace and literal **.
        else:
            for index, segment in enumerate(str(text).split("**")):
                run = paragraph.add_run(segment)
                if index % 2:
                    run.bold = True
                    if emphasis:
                        run.font.color.rgb = RGBColor.from_string("2F5496")
        return paragraph

    def _add(self, block):
        if block.kind == "paragraph":
            self._paragraph(block.args[0], block.options["style"],
                            block.options.get("emphasis", False))
        elif block.kind == "heading":
            self.doc.add_paragraph(block.args[0], "Heading " + str(block.options["level"]))
        elif block.kind == "equation":
            paragraph = self.doc.add_paragraph(style="公式")
            math = OxmlElement("m:oMath")
            _math_append(math, block.args)
            paragraph._p.append(math)
        elif block.kind == "table":
            self._table(*block.args, **block.options)
        elif block.kind == "image":
            self._image(*block.args, **block.options)
        elif block.kind == "toc":
            self.doc.add_paragraph("目录", "目录标题")
            paragraph = self.doc.add_paragraph(style="toc 1")
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
            field.set(qn("w:dirty"), "true")
            # A headless conversion may preserve the field without evaluating it.
            # Keep that state visible instead of silently presenting an empty TOC.
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "请在Word中更新目录字段以显示条目和页码"
            run.append(text)
            field.append(run)
            paragraph._p.append(field)
            settings = self.doc.settings.element
            update = settings.find(qn("w:updateFields"))
            if update is None:
                update = OxmlElement("w:updateFields")
                settings.append(update)
            update.set(qn("w:val"), "true")
        else:
            raise ValueError("Unknown content block: " + block.kind)
        self._after_table = block.kind == "table"

    def _table(self, headers, rows, widths_cm, *, allow_row_split=False):
        columns = len(headers)
        if not columns or len(widths_cm) != columns or any(len(row) != columns for row in rows):
            raise ValueError("Table headers, rows, and widths must have matching columns")
        if any(width <= 0 for width in widths_cm) or sum(widths_cm) > 16.01:
            raise ValueError("Table widths must fit the 16 cm text area")
        table = self.doc.add_table(rows=1, cols=columns)
        table.autofit = False
        for column, width in zip(table.columns, widths_cm):
            column.width = Cm(width)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement("w:" + side)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                border.set(qn("w:" + key), value)
            borders.append(border)
        table._tbl.tblPr.append(borders)
        margins = OxmlElement("w:tblCellMar")
        for side, amount in (("top", 0.18), ("bottom", 0.18), ("left", 0.2), ("right", 0.2)):
            edge = OxmlElement("w:" + side)
            edge.set(qn("w:w"), str(round(Cm(amount).twips)))
            edge.set(qn("w:type"), "dxa")
            margins.append(edge)
        table._tbl.tblPr.append(margins)
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for index, values in enumerate((headers, *rows)):
            row = table.rows[0] if index == 0 else table.add_row()
            if not allow_row_split:
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell, value, width in zip(row.cells, values, widths_cm):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                paragraph = cell.paragraphs[0]
                paragraph.style = "表格正文"
                paragraph.add_run(str(value)).bold = index == 0
                if index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EEEEEE")
                    cell._tc.get_or_add_tcPr().append(shade)

    def _image(self, path, caption, *, alt_text, width_cm):
        section = self.doc.sections[0]
        available = (section.page_width - section.left_margin - section.right_margin) / 360000
        width = available if width_cm is None else min(width_cm, available)
        if width <= 0:
            raise ValueError("Image width must be positive")
        paragraph = self.doc.add_paragraph(style="图片")
        picture = paragraph.add_run().add_picture(str(path), width=Cm(width))
        picture._inline.docPr.set("descr", alt_text or caption)
        self.doc.add_paragraph(caption, "图注")

    def _source(self, item):
        paragraph = self.doc.add_paragraph(style="来源")
        if item.url:
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), paragraph.part.relate_to(item.url, RT.HYPERLINK, is_external=True))
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:" + key), self.chinese_font)
            props.append(fonts)
            for key, value in (("sz", "19"), ("color", "000000")):
                prop = OxmlElement("w:" + key)
                prop.set(qn("w:val"), value)
                props.append(prop)
            run.append(props)
            text = OxmlElement("w:t")
            text.text = item.label
            run.append(text)
            link.append(run)
            paragraph._p.append(link)
        else:
            paragraph.add_run(item.label)
        if item.detail:
            paragraph.add_run("。" + item.detail)

    def save(self):
        """Serialize before exclusive creation; existing files are never overwritten."""
        buffer = BytesIO()
        self.doc.save(buffer)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        with self.output_path.open("xb") as output:
            output.write(buffer.getvalue())
        return self.output_path
