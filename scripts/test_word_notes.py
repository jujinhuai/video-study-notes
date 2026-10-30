"""Regression checks for editable math, TOC styles, and table pagination."""
import unittest
from pathlib import Path

from docx.oxml.ns import qn

from word_notes import (
    NotesDocument, column_vector, equation, matrix, sub, summation, sup, table, toc,
)


class WordNotesTests(unittest.TestCase):
    def document(self, blocks):
        # Construction is in memory; only NotesDocument.save() writes a file.
        return NotesDocument('公式检查', Path('公式检查.docx'), blocks=blocks).doc

    def test_column_vector_keeps_composite_values_in_one_column(self):
        value = column_vector(['x', [sub('α', '1'), '×', sub('λ', '1')]])
        rows = list(value.iter(qn('m:mr')))
        self.assertEqual([len(row.findall(qn('m:e'))) for row in rows], [1, 1])
        self.assertEqual(''.join(rows[1].itertext()), 'α1×λ1')

    def test_matrix_rejects_empty_and_ragged_rows(self):
        for rows in ([], [[]], [['a'], ['b', 'c']]):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                matrix(rows)

    def test_missing_limits_are_hidden_but_zero_is_preserved(self):
        for lower, upper, hidden in ((None, '', ('sub', 'sup')), ('j ∈ J', None, ('sup',)), (0, 'n', ())):
            with self.subTest(lower=lower, upper=upper):
                value = summation(lower, upper, sup('a', '2'))
                props = value.find(qn('m:naryPr'))
                for name in ('sub', 'sup'):
                    self.assertEqual(props.find(qn('m:' + name + 'Hide')) is not None, name in hidden)
                self.assertNotIn('None', ''.join(value.itertext()))
                if lower == 0:
                    self.assertEqual(''.join(value.find(qn('m:sub')).itertext()), '0')

    def test_math_blocks_can_be_reused_without_moving_their_xml(self):
        value = summation('j ∈ J', None, sup(sub('a', 'j'), '2'))
        block = equation(value)
        first, second = self.document([block]), self.document([block])
        self.assertEqual(len(first.element.xpath('.//m:nary')), 1)
        self.assertEqual(len(second.element.xpath('.//m:nary')), 1)
        self.assertIsNone(value.getparent())

    def test_table_rows_stay_together_and_header_repeats(self):
        doc = self.document([table(['项目', '值'], [['输出', '1000类']], [6, 10])])
        rows = doc.tables[0].rows
        self.assertTrue(all(row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is not None for row in rows))
        self.assertIsNotNone(rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')))

    def test_long_table_rows_can_explicitly_split(self):
        doc = self.document([table(['项目'], [['长文本']], [16], allow_row_split=True)])
        self.assertTrue(all(row._tr.get_or_add_trPr().find(qn('w:cantSplit')) is None for row in doc.tables[0].rows))

    def test_toc_uses_word_builtin_styles(self):
        doc = self.document([toc()])
        for level in (1, 2):
            style = doc.styles[f'toc {level}']
            self.assertTrue(style.builtin)
            self.assertEqual(style.style_id, f'TOC{level}')
        self.assertEqual(doc.paragraphs[-1].style.style_id, 'TOC1')
        self.assertIn('TOC', doc.element.xpath('.//w:fldSimple')[0].get(qn('w:instr')))


if __name__ == '__main__':
    unittest.main()
